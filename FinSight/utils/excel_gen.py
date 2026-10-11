import io
from openpyxl import Workbook

def export_expenses_excel(expenses):
    wb = Workbook()
    ws = wb.active
    ws.title = "Expenses"
    
    # Headers
    headers = ["ID", "Category", "Amount", "Date", "Description", "Type", "Receipt Filename"]
    ws.append(headers)
    
    for e in expenses:
        ws.append([
            e.id,
            e.category,
            e.amount,
            e.date.strftime("%Y-%m-%d"),
            e.description or "",
            e.type,
            e.receipt_filename or ""
        ])
        
    # Auto-adjust column widths
    for col in ws.columns:
        values = [str(cell.value or '') for cell in col]
        max_len = max(len(v) for v in values) if values else 0
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
        
    fp = io.BytesIO()
    wb.save(fp)
    return fp.getvalue()

def export_budgets_excel(budgets):
    wb = Workbook()
    ws = wb.active
    ws.title = "Budgets"
    
    headers = ["ID", "Category", "Amount Limit", "Period"]
    ws.append(headers)
    
    for b in budgets:
        ws.append([
            b.id,
            b.category,
            b.amount_limit,
            b.period
        ])
        
    for col in ws.columns:
        values = [str(cell.value or '') for cell in col]
        max_len = max(len(v) for v in values) if values else 0
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
        
    fp = io.BytesIO()
    wb.save(fp)
    return fp.getvalue()

def export_investments_excel(investments):
    wb = Workbook()
    ws = wb.active
    ws.title = "Investments"
    
    headers = ["ID", "Symbol", "Name", "Quantity", "Buy Price", "Current Price", "Asset Type", "Purchase Date"]
    ws.append(headers)
    
    for i in investments:
        ws.append([
            i.id,
            i.symbol,
            i.name,
            i.quantity,
            i.buy_price,
            i.current_price,
            i.asset_type,
            i.purchase_date.strftime("%Y-%m-%d")
        ])
        
    for col in ws.columns:
        values = [str(cell.value or '') for cell in col]
        max_len = max(len(v) for v in values) if values else 0
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
        
    fp = io.BytesIO()
    wb.save(fp)
    return fp.getvalue()

def export_goals_excel(goals):
    wb = Workbook()
    ws = wb.active
    ws.title = "Goals"
    
    headers = ["ID", "Goal Name", "Target Amount", "Current Saved Amount", "Deadline", "Category", "Completion %"]
    ws.append(headers)
    
    for g in goals:
        pct = (g.current_amount / g.target_amount) * 100 if g.target_amount > 0 else 0
        ws.append([
            g.id,
            g.name,
            g.target_amount,
            g.current_amount,
            g.deadline.strftime("%Y-%m-%d"),
            g.category,
            f"{pct:.1f}%"
        ])
        
    for col in ws.columns:
        values = [str(cell.value or '') for cell in col]
        max_len = max(len(v) for v in values) if values else 0
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
        
    fp = io.BytesIO()
    wb.save(fp)
    return fp.getvalue()
